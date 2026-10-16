from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib import colors
from reportlab.pdfbase.pdfmetrics import stringWidth

OUT='.'
NAVY=RGBColor(10,25,47); BLUE=RGBColor(47,128,237); CYAN=RGBColor(56,189,248); WHITE=RGBColor(245,248,252); MUTED=RGBColor(176,191,210); GOLD=RGBColor(250,204,21)
slides=[
('Orbit','One place for everything competing for a student\'s attention','CSE 4333 Cloud Computing | Midpoint Evaluation', 'A student responsibility dashboard that turns scattered obligations into a clear, actionable view.'),
('The problem is fragmentation','Students manage school, work, leadership, projects, and life across separate tools.','ORBIT RESPONSE', 'One dashboard keeps the responsibility, priority, due date, category, and completion state together.'),
('A working cloud application','The core loop is complete: create, persist, retrieve, filter, complete, and summarize.','VERIFIED FUNCTIONALITY', 'Responsibilities persist in DynamoDB, while the React dashboard keeps the current state visible.'),
('Architecture shaped by real constraints','AWS Academy permissions changed the design, so the team adapted without losing cloud security.','DEPLOYMENT PATH', 'React → FastAPI Proxy → IAM-protected Lambda Function URL → Lambda → DynamoDB'),
('Midpoint result and next steps','Orbit has a repeatable cloud foundation and a testable product loop.','NEXT MILESTONE', 'Polish the user experience, expand validation, and strengthen deployment and test evidence.'),
]

def add_bg(slide, idx):
    bg=slide.background.fill; bg.solid(); bg.fore_color.rgb=NAVY
    # accent blocks
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5)); s.fill.solid(); s.fill.fore_color.rgb=CYAN; s.line.fill.background()
    for x,y,w,h,c in [(10.8,0.35,1.2,0.08,BLUE),(11.35,0.55,0.65,0.08,CYAN),(0.55,7.05,2.0,0.03,BLUE)]:
        q=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=c; q.line.fill.background()
    tb=slide.shapes.add_textbox(Inches(11.8), Inches(7.08), Inches(0.7), Inches(.2)); tf=tb.text_frame; tf.text=str(idx+1); tf.paragraphs[0].font.size=Pt(10); tf.paragraphs[0].font.color.rgb=MUTED; tf.paragraphs[0].alignment=PP_ALIGN.RIGHT

def add_text(slide, text, x,y,w,h, size, color=WHITE, bold=False, font='Aptos', align=PP_ALIGN.LEFT):
    tb=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=tb.text_frame; tf.word_wrap=True; tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    p=tf.paragraphs[0]; p.text=text; p.alignment=align; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color
    return tb

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
for i,(title,sub,kicker,body) in enumerate(slides):
    sl=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(sl,i)
    add_text(sl,kicker,0.75,0.65,5,.3,11,CYAN,True)
    add_text(sl,title,0.75,1.12,11.5,0.8,30,WHITE,True)
    add_text(sl,sub,0.75,2.15,11.3,0.55,17,MUTED)
    # content card
    card=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.75), Inches(3.05), Inches(11.85), Inches(2.7)); card.fill.solid(); card.fill.fore_color.rgb=RGBColor(20,42,72); card.line.color.rgb=RGBColor(41,72,111)
    if i==0:
        add_text(sl,'Orbit',1.2,3.45,4,0.8,44,CYAN,True); add_text(sl,body,1.2,4.45,9.8,.7,20,WHITE)
    elif i==1:
        add_text(sl,'One shared view',1.2,3.48,4.5,.55,25,WHITE,True); add_text(sl,body,1.2,4.32,10.2,1.0,18,MUTED)
        for j,t in enumerate(['Academic','Work','Leadership','Personal']):
            add_text(sl,t,1.2+j*2.55,5.35,2.2,.25,11,CYAN,True)
    elif i==2:
        add_text(sl,'Create → save → return → act',1.2,3.48,8,.55,25,WHITE,True); add_text(sl,body,1.2,4.32,10.4,.8,18,MUTED)
        for j,t in enumerate(['DynamoDB persistence','Dashboard statistics','Category filters','Completion state']): add_text(sl,t,1.2+(j%2)*5.3,5.25+(j//2)*.35,4.8,.22,12,CYAN,True)
    elif i==3:
        add_text(sl,'Secure request path',1.2,3.42,5,.45,24,WHITE,True); add_text(sl,body,1.2,4.18,10.4,.7,17,MUTED)
        add_text(sl,'IAM + SigV4',1.2,5.2,2.5,.3,13,GOLD,True); add_text(sl,'SAM / CloudFormation',4.2,5.2,3.5,.3,13,GOLD,True); add_text(sl,'Repeatable deployment',8.25,5.2,3.5,.3,13,GOLD,True)
    else:
        add_text(sl,'Working proof at midpoint',1.2,3.42,6,.5,24,WHITE,True); add_text(sl,body,1.2,4.15,10.4,.65,17,MUTED)
        add_text(sl,'Next: stronger evidence, smoother UX, broader testing',1.2,5.25,10,.3,14,CYAN,True)
    add_text(sl,'ORBIT  /  CSE 4333',.75,6.8,2.5,.2,9,MUTED,True)

prs.save('Orbit_Midpoint_Evaluation.pptx')

doc=Document(); sec=doc.sections[0]; sec.top_margin=DInches(.65); sec.bottom_margin=DInches(.65); sec.left_margin=DInches(.8); sec.right_margin=DInches(.8)
styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=DPt(10.5); styles['Normal'].font.color.rgb=DRGB(38,50,65)
doc.add_heading('Orbit Midpoint Evaluation',0); p=doc.add_paragraph('Speaking notes and demonstration walkthrough'); p.style='Subtitle'
doc.add_paragraph('CSE 4333 Cloud Computing | Five-minute presentation plus five-to-ten-minute live demonstration')
for i,(title,sub,kicker,body) in enumerate(slides,1):
    doc.add_heading(f'Slide {i}: {title}',1)
    doc.add_paragraph(f'Slide message: {sub}')
    notes=[
        'Introduce Orbit as a student responsibility management dashboard. The goal is to reduce the mental overhead of switching between separate lists, calendars, and notes. Frame the project as a practical cloud application with a clear user problem.',
        'Explain that the problem is fragmentation, not a lack of productivity tools. Students move between academic deadlines, work shifts, leadership commitments, and personal tasks. Orbit gives each responsibility the same basic structure so the dashboard can show what matters now.',
        'Describe the verified product loop. A user creates a responsibility, optionally adds a due date, selects category and priority, and saves it. The application can retrieve the saved record, filter the list, mark it completed, and show summary statistics. Emphasize that the state persists in DynamoDB.',
        'Explain the architecture and the design constraint. AWS Academy permissions blocked API Gateway resources, so the team used an IAM-protected Lambda Function URL and a FastAPI proxy. SigV4 request signing protects the cloud request. SAM and CloudFormation make the infrastructure repeatable instead of manual.',
        'Close with the midpoint assessment: the team has a working cloud foundation and a testable product loop. The next milestone is product polish and stronger evidence, especially around validation, deployment repeatability, and test coverage.'
    ][i-1]
    doc.add_paragraph(notes)
doc.add_heading('Detailed demonstration walkthrough',0)
doc.add_paragraph('Target duration: 5-10 minutes. Keep the live path short and narrate what each action proves.')
steps=[
('1. Open the dashboard (30 seconds)','Start on the Orbit React dashboard. Point out the main list, category filter, summary statistics, and the add-responsibility control.'),
('2. Create a responsibility (60-90 seconds)','Add a realistic item such as “Submit cloud architecture reflection.” Select Academic, choose a high priority, and set an optional due date. Save it. Explain that the frontend sends the request through the local FastAPI proxy.'),
('3. Show persistence (60 seconds)','Refresh the page or reload the list. Confirm the responsibility remains. State that the cloud path ends in Lambda writing to DynamoDB, so the dashboard is not relying on browser-only state.'),
('4. Use filtering and statistics (60 seconds)','Filter by Academic or another category. Point out the list change and the dashboard statistics. This demonstrates that the application can organize the stored records into a useful view.'),
('5. Mark complete (45 seconds)','Mark the item completed. Show the visual state change and any updated statistics. Explain that completion status is persisted with the responsibility.'),
('6. Tie back to cloud engineering (60-120 seconds)','Briefly show the repository or deployment configuration if time allows: SAM template, Lambda configuration, and DynamoDB table definition. Mention the API Gateway permission issue, the Function URL decision, resource-based DynamoDB permissions, the changed AWS account, and the value of repeatable SAM/CloudFormation deployment.'),
('7. Close (30 seconds)','Return to the dashboard and summarize: Orbit solves a real coordination problem, the core data loop works, and the next step is improving polish and evidence.'),
]
for h,b in steps: doc.add_heading(h,2); doc.add_paragraph(b)
doc.add_heading('Presenter reminders',1); doc.add_paragraph('Keep the product story in front of the infrastructure story. Use the architecture to explain why the app is dependable and secure, not as a list of services. If the live AWS environment is unavailable, use screenshots or a recorded run and explain which persistence and deployment checks were verified previously.')
doc.save('Orbit_Presentation_Notes.docx')

# PDF preview matching the deck's visual language
c=canvas.Canvas('Orbit_Midpoint_Evaluation.pdf', pagesize=landscape(letter)); W,H=landscape(letter)
def wrap(c,text,x,y,maxw,size,leading,color):
    words=text.split(); line=''; lines=[]
    for w in words:
        test=(line+' '+w).strip()
        if stringWidth(test,'Helvetica',size)>maxw and line: lines.append(line); line=w
        else: line=test
    if line: lines.append(line)
    c.setFillColor(color); c.setFont('Helvetica',size)
    for ln in lines: c.drawString(x,y,ln); y-=leading
    return y
for i,(title,sub,kicker,body) in enumerate(slides):
    c.setFillColor(colors.HexColor('#0A192F')); c.rect(0,0,W,H,fill=1,stroke=0); c.setFillColor(colors.HexColor('#38BDF8')); c.rect(0,0,10,H,fill=1,stroke=0)
    c.setFillColor(colors.HexColor('#38BDF8')); c.setFont('Helvetica-Bold',10); c.drawString(48,H-44,kicker)
    c.setFillColor(colors.HexColor('#F5F8FC')); c.setFont('Helvetica-Bold',27); c.drawString(48,H-88,title)
    wrap(c,sub,48,H-118,680,15,20,colors.HexColor('#B0BFD2'))
    c.setFillColor(colors.HexColor('#142A48')); c.roundRect(48,90,696,190,12,fill=1,stroke=0)
    c.setFillColor(colors.HexColor('#F5F8FC')); c.setFont('Helvetica-Bold',22); c.drawString(75,225, 'Orbit' if i==0 else ('One shared view' if i==1 else ('Create -> save -> return -> act' if i==2 else ('Secure request path' if i==3 else 'Working proof at midpoint'))))
    wrap(c,body,75,185,620,15,21,colors.HexColor('#B0BFD2'))
    c.setFillColor(colors.HexColor('#B0BFD2')); c.setFont('Helvetica-Bold',8); c.drawString(48,42,'ORBIT  /  CSE 4333'); c.drawRightString(744,42,str(i+1)); c.showPage()
c.save()
print('created')
